import os
from docx import Document


def generate_docx_report(data):

    os.makedirs(
        "generated_reports",
        exist_ok=True
    )

    path = (
        "generated_reports/"
        "maintenance_report.docx"
    )

    document = Document()

    document.add_heading(
        "AI Software Maintenance Report",
        0
    )

    for key, value in data.items():

        document.add_heading(
            str(key),
            level=2
        )

        document.add_paragraph(
            str(value)
        )

    document.save(path)

    return path